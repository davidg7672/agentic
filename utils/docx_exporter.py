"""
Generate .docx files from the same LaTeX source that produces the PDF,
so both exports look identical and the DOCX is fully editable.
"""
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

_MARGIN_IN = 0.65       # left/right inches — matches LaTeX geometry
_TOP_BOT_IN = 0.50      # top/bottom inches
_TEXT_W = 8.5 - 2 * _MARGIN_IN  # 7.2 inches usable width
_FONT = "Calibri"
_BODY_PT = Pt(10)


# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------

def _init_doc() -> Document:
    doc = Document()
    # Explicit baseline so all paragraphs start from a known font/size
    normal = doc.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = _BODY_PT
    for section in doc.sections:
        section.top_margin = Inches(_TOP_BOT_IN)
        section.bottom_margin = Inches(_TOP_BOT_IN)
        section.left_margin = Inches(_MARGIN_IN)
        section.right_margin = Inches(_MARGIN_IN)
    # Remove the default blank paragraph Word always adds
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    return doc


def _zero(para) -> None:
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)


def _bottom_border(para) -> None:
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "000000")
    pBdr.append(bot)
    pPr.append(pBdr)


def _right_tab(para) -> None:
    pPr = para._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(_TEXT_W * 1440)))
    tabs.append(tab)
    pPr.append(tabs)


# ---------------------------------------------------------------------------
# LaTeX → plain text helpers
# ---------------------------------------------------------------------------

def _delatex(text: str) -> str:
    """Strip LaTeX markup and unescape special characters."""
    text = re.sub(r"\\href\{[^}]*\}\{([^}]*)\}", r"\1", text)   # \href{url}{label}
    text = re.sub(r"\\text(?:bf|it|rm|tt|sc|sf)\{([^}]*)\}", r"\1", text)  # \textXX{}
    text = re.sub(r"\\(?:small|large|Large|LARGE|normalsize|footnotesize|hfill)\b\s*", "", text)
    text = text.replace(r"---", "—").replace(r"--", "–")
    text = text.replace(r"\&", "&").replace(r"\$", "$").replace(r"\%", "%")
    text = text.replace(r"\#", "#").replace(r"\_", "_")
    text = text.replace(r"\{", "{").replace(r"\}", "}")
    text = text.replace(r"\textasciitilde{}", "~").replace(r"\textasciicircum{}", "^")
    text = text.replace(r"\textbackslash{}", "\\")
    text = text.rstrip("\\").strip()  # trailing LaTeX line-break \\
    return text.strip()


def _add_runs(para, text: str, size: Pt, bold: bool = False, italic: bool = False) -> None:
    """Parse \\textbf{} and \\textit{} in *text* and add styled runs to *para*."""
    pattern = re.compile(
        r"\\textbf\{([^}]*)\}"
        r"|\\textit\{([^}]*)\}"
        r"|\\href\{[^}]*\}\{([^}]*)\}"
        r"|([^\\]+|.)"
    )
    for m in pattern.finditer(text):
        if m.group(1) is not None:
            r = para.add_run(_delatex(m.group(1)))
            r.bold = True
            r.italic = italic
            r.font.size = size
        elif m.group(2) is not None:
            r = para.add_run(_delatex(m.group(2)))
            r.bold = bold
            r.italic = True
            r.font.size = size
        elif m.group(3) is not None:
            r = para.add_run(m.group(3))
            r.bold = bold
            r.italic = italic
            r.font.size = size
        else:
            chunk = m.group(4)
            if chunk and not chunk.startswith("\\"):
                r = para.add_run(_delatex(chunk))
                r.bold = bold
                r.italic = italic
                r.font.size = size


# ---------------------------------------------------------------------------
# Per-element renderers
# ---------------------------------------------------------------------------

def _render_name(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _zero(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(17)    # \LARGE at 10pt base ≈ 17.28pt


def _render_contact(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _zero(p)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(_delatex(text))
    run.font.size = Pt(9)     # \small at 10pt base = 9pt


def _render_section(doc: Document, name: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)   # matches \titlespacing* before
    p.paragraph_format.space_after = Pt(4)    # matches \titlespacing* after
    run = p.add_run(name)
    run.bold = True
    run.font.size = _BODY_PT
    _bottom_border(p)


def _render_entry_header(doc: Document, left: str, right: str) -> None:
    """Bold company/school on the left, plain location on the right."""
    p = doc.add_paragraph()
    _zero(p)
    p.paragraph_format.space_before = Pt(3)   # small gap above each new entry
    _right_tab(p)
    r1 = p.add_run(_delatex(left))
    r1.bold = True
    r1.font.size = _BODY_PT
    tab = p.add_run("\t")
    tab.font.size = _BODY_PT
    r2 = p.add_run(_delatex(right))
    r2.font.size = _BODY_PT


def _render_entry_sub(doc: Document, left: str, right: str) -> None:
    """Italic title on the left, plain date on the right."""
    p = doc.add_paragraph()
    _zero(p)
    _right_tab(p)
    r1 = p.add_run(_delatex(left))
    r1.italic = True
    r1.font.size = _BODY_PT
    tab = p.add_run("\t")
    tab.font.size = _BODY_PT
    r2 = p.add_run(_delatex(right))
    r2.font.size = _BODY_PT


def _render_skillrow(doc: Document, label: str, skills: str) -> None:
    """Bold label + skills — mirrors \\skillrow in LaTeX. Long lists wrap naturally."""
    p = doc.add_paragraph()
    _zero(p)
    p.paragraph_format.space_after = Pt(1)
    r1 = p.add_run(_delatex(label) + " ")
    r1.bold = True
    r1.font.size = _BODY_PT
    r2 = p.add_run(_delatex(skills))
    r2.font.size = _BODY_PT


def _render_bullet(doc: Document, text: str) -> None:
    """Bullet point using a plain paragraph + unicode bullet character.

    Avoids Word's 'List Bullet' style whose built-in numPr XML overrides any
    left_indent we set, producing inconsistent indentation across platforms.
    """
    p = doc.add_paragraph()
    _zero(p)
    p.paragraph_format.left_indent = Pt(15)    # 1.5em at 10pt — matches LaTeX leftmargin
    p.paragraph_format.first_line_indent = Pt(-10)  # bullet hangs left of text
    p.paragraph_format.space_after = Pt(1)
    r_sym = p.add_run("• ")               # • + space
    r_sym.font.size = _BODY_PT
    _add_runs(p, text, _BODY_PT)


def _render_bold_line(doc: Document, text: str) -> None:
    """Standalone bold line (project header, key:value, etc.)."""
    p = doc.add_paragraph()
    _zero(p)
    p.paragraph_format.space_before = Pt(3)    # matches entry_header gap
    p.paragraph_format.space_after = Pt(1)
    _add_runs(p, text, _BODY_PT)


# ---------------------------------------------------------------------------
# Main entry points
# ---------------------------------------------------------------------------

def resume_to_docx(latex_source: str) -> bytes:
    """
    Parse the LaTeX resume source (same string used for PDF) and produce a
    visually matching, fully-editable DOCX.
    """
    doc = _init_doc()

    # Extract body between \begin{document} … \end{document}
    body_m = re.search(r"\\begin\{document\}(.*?)\\end\{document\}", latex_source, re.DOTALL)
    body = body_m.group(1) if body_m else latex_source

    lines = body.splitlines()
    i = 0
    in_center = False
    in_itemize = False
    center_buf: list[str] = []

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        i += 1

        # ── Skip blanks and comments ──────────────────────────────────────
        if not line or line.startswith("%"):
            continue

        # ── Center environment ────────────────────────────────────────────
        if line == r"\begin{center}":
            in_center = True
            center_buf = []
            continue

        if line == r"\end{center}":
            in_center = False
            _flush_center(doc, center_buf)
            continue

        if in_center:
            center_buf.append(line)
            continue

        # ── Itemize environment ───────────────────────────────────────────
        if line == r"\begin{itemize}":
            in_itemize = True
            continue

        if line == r"\end{itemize}":
            in_itemize = False
            continue

        if in_itemize:
            m = re.match(r"\\item\s+(.*)", line)
            if m:
                _render_bullet(doc, m.group(1))
            continue

        # ── Section header ────────────────────────────────────────────────
        m = re.match(r"\\section\*\{([^}]+)\}", line)
        if m:
            _render_section(doc, m.group(1))
            continue

        # ── \skillrow{Label:}{skills} ─────────────────────────────────────
        m = re.match(r"\\skillrow\{([^}]+)\}\{([^}]+)\}", line)
        if m:
            _render_skillrow(doc, m.group(1), m.group(2))
            continue

        # ── Entry header: \textbf{Company}\hfill Location\\ ───────────────
        m = re.match(r"\\textbf\{([^}]+)\}\\hfill\s+(.+)", line)
        if m:
            _render_entry_header(doc, m.group(1), m.group(2))
            continue

        # ── Entry subtitle: \textit{Title}\hfill Date ─────────────────────
        m = re.match(r"\\textit\{([^}]+)\}\\hfill\s+(.+)", line)
        if m:
            _render_entry_sub(doc, m.group(1), m.group(2))
            continue

        # ── Standalone \textbf{...} (project header or key:value line) ────
        if line.startswith(r"\textbf{") or "\\textbf{" in line:
            _render_bold_line(doc, line)
            continue

        # ── Generic text ──────────────────────────────────────────────────
        cleaned = _delatex(line)
        if cleaned:
            p = doc.add_paragraph()
            _zero(p)
            p.paragraph_format.space_after = Pt(1)
            p.add_run(cleaned).font.size = _BODY_PT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _flush_center(doc: Document, lines: list[str]) -> None:
    """Render the buffered center-environment lines (name + contact)."""
    for line in lines:
        # Name: {\LARGE\textbf{Name}}\\[2pt]  or  {\LARGE\textbf{Name}}
        m = re.match(r"\{\\LARGE\\textbf\{([^}]+)\}\}", line)
        if m:
            _render_name(doc, m.group(1))
            continue

        # Contact line (starts with \small or contains @ or |)
        cleaned = _delatex(line)
        if cleaned:
            _render_contact(doc, cleaned)


def cover_letter_to_docx(letter_text: str) -> bytes:
    """Convert plain-text (or lightly markdown-formatted) cover letter to DOCX."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    # Strip markdown bold/italic before rendering
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", letter_text)
    text = re.sub(r"\*([^*]+?)\*", r"\1", text)

    for block in text.strip().split("\n\n"):
        block = block.strip()
        if not block:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(block)
        run.font.size = Pt(10.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
